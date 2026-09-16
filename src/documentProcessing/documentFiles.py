"""One class per file format, each collecting that format's metadata.

`buildDocumentFile` picks the class from the file's extension, so a caller
processes a pdf, a docx and a markdown file the same way and each format
keeps its own extraction in one place.

What every format has in common - lines, words, characters, tokens - is
counted once on `DocumentFile` from the text a subclass hands back, so the
counts mean the same thing whatever the file was. What only one format has -
pages in a pdf, paragraphs in a docx - goes in `DocumentMetadata.extra`,
which is also where anything added later belongs.

A file that cannot be read is not an error: its metadata comes back with the
counts left as None and `notes` saying why, so one unreadable file inside an
archive of a hundred does not lose the other ninety nine.
"""

import re
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Any

import docx
import pdfplumber
import tiktoken

from documentProcessing.downloader import LoadedFile

# The encoding GPT-4o and friends count in, so a token here is a token there.
TOKEN_ENCODING_NAME = "o200k_base"

TOKEN_METHOD_ENCODED = "tiktoken"
TOKEN_METHOD_ESTIMATED = "whitespace"

# A markdown table is a row of pipes over a row of dashes; this matches the
# dashed line, of which there is exactly one per table.
MARKDOWN_TABLE_PATTERN = re.compile(r"^\s*\|?[\s:|-]*-[-\s:|]*\|[\s:|-]*$")

TEXT_READ_ERRORS = "replace"


@dataclass(slots=True)
class DocumentMetadata:
    """What one file in a submitted document is made of."""

    fileName: str
    fileFormat: str
    sizeBytes: int
    lineCount: int | None = None
    wordCount: int | None = None
    characterCount: int | None = None
    tokenCount: int | None = None
    tableCount: int | None = None
    notes: str = ""
    # Format specific metadata, and the place for anything added later.
    extra: dict[str, Any] = field(default_factory=dict)


@lru_cache(maxsize=1)
def getTokenEncoding() -> tiktoken.Encoding | None:
    """The tokeniser, or None when it cannot be loaded.

    `tiktoken` fetches the encoding the first time it is asked for and caches
    it on disk. A machine with no way out to the internet therefore has no
    encoding to load, which is worth degrading over rather than failing a
    whole job for - `countTokens` estimates instead.
    """
    try:
        return tiktoken.get_encoding(TOKEN_ENCODING_NAME)
    except Exception:
        return None


def countTokens(text: str) -> tuple[int, str]:
    """How many tokens `text` holds, and how that number was arrived at."""
    encoding = getTokenEncoding()
    if encoding is None:
        return len(text.split()), TOKEN_METHOD_ESTIMATED
    return len(encoding.encode(text)), TOKEN_METHOD_ENCODED


def countMarkdownTables(text: str) -> int:
    """How many markdown tables `text` holds, by their separator rows."""
    return sum(
        1 for line in text.splitlines() if MARKDOWN_TABLE_PATTERN.match(line)
    )


class DocumentFile:
    """One file on disk, and the metadata that can be read out of it.

    Subclasses say what their format's text is (`readText`) and how many
    tables it holds (`countTables`); everything counted from plain text is
    handled here so every format counts it the same way.
    """

    fileFormats: tuple[str, ...] = ()

    def __init__(self, loadedFile: LoadedFile) -> None:
        """Hold the file to read; nothing is parsed until asked for."""
        self.loadedFile = loadedFile

    def readText(self) -> str:
        """All the text in the file, as one string."""
        raise NotImplementedError

    def countTables(self) -> int | None:
        """How many tables the file holds, or None when it cannot be known."""
        return None

    def collectExtra(self) -> dict[str, Any]:
        """Metadata only this format has. Empty unless a subclass adds to it."""
        return {}

    def collectMetadata(self) -> DocumentMetadata:
        """Everything known about this file.

        Blocking work - parsing a large pdf is neither quick nor async - so
        call it through `asyncio.to_thread` from a job.
        """
        metadata = DocumentMetadata(
            fileName=self.loadedFile.fileName,
            fileFormat=self.loadedFile.fileFormat,
            sizeBytes=self.loadedFile.sizeBytes,
        )

        try:
            text = self.readText()
            metadata.tableCount = self.countTables()
            metadata.extra = self.collectExtra()
        except Exception as error:
            metadata.notes = f"Content could not be read: {error}"
            return metadata

        tokenCount, tokenMethod = countTokens(text)
        metadata.lineCount = len(text.splitlines())
        metadata.wordCount = len(text.split())
        metadata.characterCount = len(text)
        metadata.tokenCount = tokenCount
        metadata.extra["tokenCountMethod"] = tokenMethod
        return metadata


class TextDocumentFile(DocumentFile):
    """A plain text file, read as it stands."""

    fileFormats = (".txt",)

    def readText(self) -> str:
        """The file's text, undecodable bytes replaced rather than fatal."""
        return self.loadedFile.data.decode("utf-8", errors=TEXT_READ_ERRORS)

    def countTables(self) -> int | None:
        """Plain text has no table structure to count."""
        return 0


class MarkdownDocumentFile(TextDocumentFile):
    """A markdown file: plain text, plus tables written in pipe syntax."""

    fileFormats = (".md",)

    def countTables(self) -> int | None:
        """Tables counted by their `|---|---|` separator rows."""
        return countMarkdownTables(self.readText())


class PdfDocumentFile(DocumentFile):
    """A pdf, read page by page for both its text and its tables."""

    fileFormats = (".pdf",)

    def readText(self) -> str:
        """Every page's text, joined in page order."""
        with pdfplumber.open(self.loadedFile.openStream()) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)

    def countTables(self) -> int | None:
        """Tables found across every page."""
        with pdfplumber.open(self.loadedFile.openStream()) as pdf:
            return sum(len(page.find_tables()) for page in pdf.pages)

    def collectExtra(self) -> dict[str, Any]:
        """How many pages the pdf has."""
        with pdfplumber.open(self.loadedFile.openStream()) as pdf:
            return {"pageCount": len(pdf.pages)}


class WordDocumentFile(DocumentFile):
    """A Word document, read as paragraphs plus its tables.

    `.doc` is the old binary format, which `python-docx` cannot open - such a
    file comes back with its counts left as None and a note saying so, unless
    it is really a `.docx` that was named `.doc`, which does open.
    """

    fileFormats = (".docx", ".doc")

    def readText(self) -> str:
        """Every paragraph, one per line, in document order."""
        document = docx.Document(self.loadedFile.openStream())
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def countTables(self) -> int | None:
        """Tables the document declares, which Word records structurally."""
        return len(docx.Document(self.loadedFile.openStream()).tables)

    def collectExtra(self) -> dict[str, Any]:
        """How many paragraphs and sections the document has."""
        document = docx.Document(self.loadedFile.openStream())
        return {
            "paragraphCount": len(document.paragraphs),
            "sectionCount": len(document.sections),
        }


class UnsupportedDocumentFile(DocumentFile):
    """A file whose format has no reader here, recorded rather than dropped.

    Archives carry stray files - images, `.DS_Store`, a stray binary - and
    knowing one was there and skipped is more useful than silence.
    """

    def readText(self) -> str:
        """Always refuses: there is no reader for this format."""
        raise ValueError(
            f"No reader for {self.loadedFile.fileFormat or 'this'} files"
        )


DOCUMENT_FILE_CLASSES: tuple[type[DocumentFile], ...] = (
    TextDocumentFile,
    MarkdownDocumentFile,
    PdfDocumentFile,
    WordDocumentFile,
)


def buildDocumentFile(loadedFile: LoadedFile) -> DocumentFile:
    """The document class that reads `loadedFile`, by its extension."""
    for documentClass in DOCUMENT_FILE_CLASSES:
        if loadedFile.fileFormat in documentClass.fileFormats:
            return documentClass(loadedFile)
    return UnsupportedDocumentFile(loadedFile)


def collectMetadata(loadedFiles: list[LoadedFile]) -> list[DocumentMetadata]:
    """Metadata for every file given, in the order given.

    Blocking work - call it through `asyncio.to_thread` from a job.
    """
    return [
        buildDocumentFile(loadedFile).collectMetadata()
        for loadedFile in loadedFiles
    ]
