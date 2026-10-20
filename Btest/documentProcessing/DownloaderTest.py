"""Btest for `documentProcessing.downloader`, on dummy files.

`dummyFolder` writes a dummy file for every accepted format - plus files
that are not what their extension says, one too big and one of a format
not taken - and deletes the folder after the module. `fileServer` hands
them out over HTTP from 127.0.0.1, so `downloadDocument` really streams
each one, with no internet involved. Every download goes to a fresh
folder that is deleted after its test.

Runs the same on Windows and Linux: folders come from `tempfile`, the
server listens on loopback only, and each file is asked for by the exact
name it was written with, since Linux file names are case-sensitive.
There is no writer for rar in Python, so the rar dummy is the real RAR5
signature followed by padding - which is all the downloader checks.

    uv run pytest Btest/documentProcessing/DownloaderTest.py -q
"""

import asyncio
import functools
import shutil
import socket
import tarfile
import tempfile
import threading
import uuid
import zipfile
from collections.abc import Iterator
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

import docx
import pytest

from documentProcessing import downloader
from documentProcessing.downloader import DownloadError

JOB_ID = str(uuid.uuid4())

# Accepted dummy files, by name, and the extension each is saved with.
GOOD_FILES = {
    "report.pdf": ".pdf",
    "SCAN.PDF": ".pdf",
    "letter.docx": ".docx",
    "notes.txt": ".txt",
    "readme.md": ".md",
    "archive.rar": ".rar",
    "bundle.zip": ".zip",
    "bundle.tar": ".tar",
}
# Files that are not what their extension says.
MISLABELED_FILES = {
    "fake.pdf": ".pdf",
    "fake.docx": ".docx",
    "fake.zip": ".zip",
    "fake.rar": ".rar",
    "fake.tar": ".tar",
    "binary.txt": ".txt",
    "binary.md": ".md",
}
TOO_BIG_FILE = "big.txt"
TOO_BIG_BYTES = 2048
# A size limit the big file is over.
SMALL_LIMIT_BYTES = 1024

# The paths the dummy server was asked for, in order.
REQUESTED_PATHS: list[str] = []
# Under this path the server sends a file without saying how big it is.
UNDECLARED_PREFIX = "/undeclared/"


def buildPdf() -> bytes:
    """The smallest well-formed pdf: a catalog and a trailer."""
    return (
        b"%PDF-1.4\n1 0 obj << /Type /Catalog >> endobj\n"
        b"trailer << /Root 1 0 R >>\n%%EOF\n"
    )


def buildDocx() -> bytes:
    """A Word document with one paragraph."""
    wordDocument = docx.Document()
    wordDocument.add_paragraph("Orders ship within two days.")
    stream = BytesIO()
    wordDocument.save(stream)
    return stream.getvalue()


def buildZip() -> bytes:
    """A zip holding one text file."""
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("inside.txt", "Inside the zip.")
    return stream.getvalue()


def buildTar() -> bytes:
    """A tar holding one text file."""
    stream = BytesIO()
    content = b"Inside the tar."
    with tarfile.open(fileobj=stream, mode="w") as archive:
        member = tarfile.TarInfo("inside.txt")
        member.size = len(content)
        archive.addfile(member, BytesIO(content))
    return stream.getvalue()


def buildRar() -> bytes:
    """The RAR5 signature, padded out like an archive's first block."""
    return b"Rar!\x1a\x07\x01\x00" + bytes(64)


def buildDummyFiles() -> dict[str, bytes]:
    """Every dummy file the server hands out, by name."""
    return {
        "report.pdf": buildPdf(),
        "SCAN.PDF": buildPdf(),
        "letter.docx": buildDocx(),
        "notes.txt": b"Delivery takes two days.\n",
        "readme.md": b"# Readme\n\nSome text.\n",
        "archive.rar": buildRar(),
        "bundle.zip": buildZip(),
        "bundle.tar": buildTar(),
        "fake.pdf": b"<html>Not a pdf.</html>",
        "fake.docx": b"Not a Word document.",
        "fake.zip": b"Not a zip.",
        "fake.rar": b"Not a rar.",
        "fake.tar": b"Not a tar.",
        "binary.txt": b"Starts as text\x00then is binary.",
        "binary.md": b"# Title\x00\x01\x02",
        TOO_BIG_FILE: b"a" * TOO_BIG_BYTES,
        "setup.exe": b"MZ not taken",
    }


class QuietFileHandler(SimpleHTTPRequestHandler):
    """Serves the dummy folder, noting each path and logging nothing."""

    def do_GET(self) -> None:
        REQUESTED_PATHS.append(self.path)
        if self.path.startswith(UNDECLARED_PREFIX):
            self.sendUndeclared(self.path.removeprefix(UNDECLARED_PREFIX))
            return
        super().do_GET()

    def sendUndeclared(self, fileName: str) -> None:
        """Send a file with no Content-Length; closing marks its end."""
        content = (Path(self.directory) / fileName).read_bytes()
        self.send_response(200)
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, format: str, *args) -> None:
        pass


@pytest.fixture(scope="module")
def dummyFolder() -> Iterator[Path]:
    """The dummy files on disk, deleted after the module's last test."""
    folder = Path(tempfile.mkdtemp(prefix="ragDummyFiles"))
    try:
        for fileName, content in buildDummyFiles().items():
            (folder / fileName).write_bytes(content)
        yield folder
    finally:
        shutil.rmtree(folder, ignore_errors=True)


@pytest.fixture(scope="module")
def fileServer(dummyFolder: Path) -> Iterator[str]:
    """The address of a local server handing out the dummy files."""
    handler = functools.partial(QuietFileHandler, directory=str(dummyFolder))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}"
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


@pytest.fixture(autouse=True)
def downloadRoot(monkeypatch) -> Iterator[Path]:
    """Downloads go to a fresh folder, deleted after the test."""
    root = Path(tempfile.mkdtemp(prefix="ragDownloads"))
    monkeypatch.setattr(downloader, "DOWNLOAD_ROOT", root)
    try:
        yield root
    finally:
        shutil.rmtree(root, ignore_errors=True)


def download(documentLink: str) -> Path:
    """Run one download for JOB_ID to the end."""
    return asyncio.run(downloader.downloadDocument(documentLink, JOB_ID))


def closedPortLink() -> str:
    """A link to a loopback port nothing is listening on."""
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    return f"http://127.0.0.1:{port}/report.pdf"


class TestCheckDocumentLink:
    """Only http(s) links to an accepted format are taken."""

    @pytest.mark.parametrize(
        "documentLink",
        [
            "https://files.example/report.pdf",
            "https://files.example/SCAN.PDF",
            "http://files.example/letter.docx",
            "https://files.example/notes.txt",
            "https://files.example/readme.md",
            "https://files.example/archive.rar",
            "https://files.example/bundle.zip",
            "https://files.example/bundle.tar",
            "https://files.example/a/b/report.pdf?signature=abc",
        ],
    )
    def testTakesAcceptedLinks(self, documentLink: str) -> None:
        assert downloader.checkDocumentLink(documentLink) is None

    @pytest.mark.parametrize(
        "documentLink",
        [
            "https://files.example/setup.exe",
            "https://files.example/old.doc",
            "https://files.example/bundle.tar.gz",
            "https://files.example/report",
            "https://files.example/",
            "ftp://files.example/report.pdf",
            "files.example/report.pdf",
            "https:///report.pdf",
        ],
    )
    def testRefusesOtherLinks(self, documentLink: str) -> None:
        assert downloader.checkDocumentLink(documentLink) is not None

    def testRefusesMissingLink(self) -> None:
        assert downloader.checkDocumentLink(None) is not None


class TestDownloadDocument:
    """A dummy file is streamed to the job's folder, byte for byte."""

    @pytest.mark.parametrize("fileName", list(GOOD_FILES))
    def testDownloadsEachAcceptedFormat(
        self,
        fileServer: str,
        dummyFolder: Path,
        downloadRoot: Path,
        fileName: str,
    ) -> None:
        path = download(f"{fileServer}/{fileName}")

        assert path == downloadRoot / JOB_ID / f"document{GOOD_FILES[fileName]}"
        assert path.read_bytes() == (dummyFolder / fileName).read_bytes()

    @pytest.mark.parametrize("fileName", list(MISLABELED_FILES))
    def testFileThatIsNotItsFormatIsRefused(
        self, fileServer: str, downloadRoot: Path, fileName: str
    ) -> None:
        fileFormat = MISLABELED_FILES[fileName]

        with pytest.raises(DownloadError, match=f"not a \\{fileFormat} file"):
            download(f"{fileServer}/{fileName}")

        assert not (downloadRoot / JOB_ID).exists()

    def testFormatNotTakenIsRefusedBeforeDownloading(
        self, fileServer: str, downloadRoot: Path
    ) -> None:
        REQUESTED_PATHS.clear()

        with pytest.raises(DownloadError, match="is not one of"):
            download(f"{fileServer}/setup.exe")

        assert REQUESTED_PATHS == []
        assert not (downloadRoot / JOB_ID).exists()

    def testMissingFileFails(
        self, fileServer: str, downloadRoot: Path
    ) -> None:
        with pytest.raises(DownloadError, match="HTTP 404"):
            download(f"{fileServer}/missing.pdf")

        assert not (downloadRoot / JOB_ID).exists()

    def testDeclaredSizeOverTheLimitIsRefusedUnread(
        self, fileServer: str, downloadRoot: Path, monkeypatch
    ) -> None:
        monkeypatch.setattr(downloader, "MAX_DOCUMENT_BYTES", SMALL_LIMIT_BYTES)
        realWritePieces = downloader.writePieces
        writes = []

        async def recordingWritePieces(response, path: Path) -> None:
            writes.append(path)
            await realWritePieces(response, path)

        monkeypatch.setattr(downloader, "writePieces", recordingWritePieces)

        with pytest.raises(DownloadError, match="over the 1024 byte limit"):
            download(f"{fileServer}/{TOO_BIG_FILE}")

        assert writes == []
        assert not (downloadRoot / JOB_ID).exists()

    def testUndeclaredSizeOverTheLimitIsCutOff(
        self, fileServer: str, downloadRoot: Path, monkeypatch
    ) -> None:
        monkeypatch.setattr(downloader, "MAX_DOCUMENT_BYTES", SMALL_LIMIT_BYTES)

        with pytest.raises(DownloadError, match="over the 1024 byte limit"):
            download(f"{fileServer}{UNDECLARED_PREFIX}{TOO_BIG_FILE}")

        assert not (downloadRoot / JOB_ID).exists()

    def testUndeclaredSizeIsReadToTheEnd(
        self, fileServer: str, dummyFolder: Path
    ) -> None:
        path = download(f"{fileServer}{UNDECLARED_PREFIX}notes.txt")

        assert path.read_bytes() == (dummyFolder / "notes.txt").read_bytes()

    def testUnreachableServerFails(self, downloadRoot: Path) -> None:
        with pytest.raises(DownloadError, match="could not be downloaded"):
            download(closedPortLink())

        assert not (downloadRoot / JOB_ID).exists()


class TestRemoveDownload:
    """A job's download folder is deleted whole, or not there to delete."""

    def testRemovesTheJobFolder(self, downloadRoot: Path) -> None:
        folder = downloadRoot / JOB_ID
        folder.mkdir()
        (folder / "document.pdf").write_bytes(buildPdf())

        downloader.removeDownload(JOB_ID)

        assert not folder.exists()

    def testMissingFolderIsNotAnError(self, downloadRoot: Path) -> None:
        downloader.removeDownload(JOB_ID)

        assert not (downloadRoot / JOB_ID).exists()
